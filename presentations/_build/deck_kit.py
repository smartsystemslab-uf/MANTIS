"""Tiny layout toolkit over python-pptx with a real text-fit checker.

Every text box is measured with Arial metrics (PIL) at build time; anything
whose wrapped height exceeds its box is recorded in WARNINGS, so overflow is
caught without needing a renderer.
"""
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import os

FONT = "Arial"


def _first_existing(*paths):
    for p in paths:
        if os.path.exists(p):
            return p
    raise FileNotFoundError(
        "No metric-compatible font found for the text-fit checker. Install Arial "
        "(macOS ships it) or Liberation Sans/Mono, or extend _first_existing() here."
    )


# Real font metrics drive the overflow checker, so a deck can be verified
# without a renderer. macOS paths first, then common Linux (Liberation).
_REG = _first_existing("/System/Library/Fonts/Supplemental/Arial.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf")
_BOLD = _first_existing("/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")
_MONO = _first_existing("/System/Library/Fonts/Supplemental/Courier New.ttf", "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf")
_cache = {}

NAVY, INK, SLATE = "0B1F3A", "12263F", "51627A"
TEAL, AMBER, RED, GREEN = "0E9AA7", "F2A900", "D64545", "2E9E6B"
PAPER, TINT, LINE, WHITE = "FFFFFF", "EAF0F6", "CBD6E2", "FFFFFF"
MIST = "9FB3C8"

WARNINGS = []


def rgb(h):
    return RGBColor.from_string(h)


def _font(path, size_pt):
    key = (path, size_pt)
    if key not in _cache:
        # 10 px per pt gives sub-pixel-accurate widths
        _cache[key] = ImageFont.truetype(path, int(size_pt * 10))
    return _cache[key]


def _wrap_lines(text, path, size_pt, width_in):
    f = _font(path, size_pt)
    max_px = width_in * 72 * 10
    total = 0
    for para in text.split("\n"):
        words = para.split(" ")
        line = ""
        n = 1
        for w in words:
            trial = (line + " " + w).strip()
            if f.getlength(trial) <= max_px or not line:
                line = trial
            else:
                n += 1
                line = w
        total += n
    return total


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.n = 0

    def slide(self, bg=PAPER, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        s._idx = self.n
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(bg)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, line_w=0.75):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    sp.text_frame.text = ""
    return sp


def card(slide, x, y, w, h, fill=TINT, line=None):
    return rect(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align="l", anchor="t",
         font=FONT, line_spacing=1.08, space_after=0, margin=0.0, name=None):
    """paras: str | list of str | list of dict(text,size,color,bold,font,space_after,runs=[(text,{...})])."""
    if isinstance(paras, str):
        paras = [paras]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    est_h = 0.0
    inner_w = w - 2 * margin
    for i, p in enumerate(paras):
        if isinstance(p, str):
            p = {"text": p}
        psize = p.get("size", size)
        pbold = p.get("bold", bold)
        pfont = p.get("font", font)
        pcolor = p.get("color", color)
        pafter = p.get("space_after", space_after)
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p.get("align", align)]
        para.line_spacing = line_spacing
        para.space_after = Pt(pafter)
        runs = p.get("runs") or [(p["text"], {})]
        full = ""
        widest_bold = pbold
        for rt, ro in runs:
            r = para.add_run()
            r.text = rt
            r.font.name = ro.get("font", pfont)
            r.font.size = Pt(ro.get("size", psize))
            r.font.bold = ro.get("bold", pbold)
            r.font.color.rgb = rgb(ro.get("color", pcolor))
            full += rt
            widest_bold = widest_bold or ro.get("bold", False)
        path = _MONO if pfont == "Courier New" else (_BOLD if widest_bold else _REG)
        lines = _wrap_lines(full, path, psize, inner_w)
        est_h += lines * psize * 1.2 * line_spacing / 72 + pafter / 72
    if est_h > h + 0.02:
        WARNINGS.append(f"slide {slide._idx}: text overflows ({est_h:.2f}in > {h:.2f}in): {str(paras[0])[:60]}")
    return box


def badge(slide, x, y, d, label, fill=TEAL, color=WHITE, size=14):
    c = rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(color)
    return c


def chip(slide, x, y, w, h, label, fill, color=WHITE, size=11):
    c = rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(color)
    if _wrap_lines(label, _BOLD, size, w - 0.12) * size * 1.2 / 72 > h + 0.02:
        WARNINGS.append(f"slide {slide._idx}: chip text overflows: {label}")
    return c


def arrow(slide, x, y, w, h, fill=MIST):
    return rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.RIGHT_ARROW)


def title(slide, s, dark=False, sub=None):
    text(slide, 0.6, 0.45, 12.1, 0.7, s, size=30, bold=True, color=WHITE if dark else NAVY, anchor="t")
    if sub:
        text(slide, 0.6, 1.2, 12.1, 0.5, sub, size=16, color=MIST if dark else SLATE)


def table(slide, x, y, col_w, rows, header, row_h=0.55, size=13, head_fill=NAVY, zebra=TINT,
          cell_colors=None):
    """Native pptx table. rows: list of lists of str. cell_colors: {(r,c): hex} text colors."""
    nrows, ncols = len(rows) + 1, len(col_w)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for c, w in enumerate(col_w):
        tbl.columns[c].width = Inches(w)
    for r in range(nrows):
        tbl.rows[r].height = Inches(row_h)
    def fillcell(cell, txt, bold, color, fill):
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r_ = p.add_run()
        r_.text = txt
        r_.font.name = FONT
        r_.font.size = Pt(size)
        r_.font.bold = bold
        r_.font.color.rgb = rgb(color)
    for c, h in enumerate(header):
        fillcell(tbl.cell(0, c), h, True, WHITE, head_fill)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            color = (cell_colors or {}).get((r - 1, c), INK)
            bold = (r - 1, c) in (cell_colors or {})
            fillcell(tbl.cell(r, c), val, bold, color, zebra if r % 2 == 0 else PAPER)
            lines = _wrap_lines(val, _BOLD if bold else _REG, size, col_w[c] - 0.2)
            if lines * size * 1.2 / 72 + 0.08 > row_h + 0.02:
                WARNINGS.append(f"slide {slide._idx}: table cell overflows row height: {val[:50]}")
    return gf
